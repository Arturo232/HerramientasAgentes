# -*- coding: utf-8 -*-
"""Utilidades para DOCX: vinetas reales, relleno de plantillas y export a PDF.

Muchas plantillas academicas NO traen el estilo de lista, por lo que las
vinetas terminan como texto "•" y al pasar a Google Docs no quedan como listas.
Este modulo inyecta numeracion real (abstractNum/num) para que Word/Google Docs
las reconozcan.

Uso como libreria:
    import docx_util
    docx_util.build(plantilla, salida, encabezado, secciones)

Uso por linea de comandos:
    python docx_util.py --plantilla P.docx --contenido c.json --salida out.docx [--pdf]
"""
import json
import os
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

_RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
if _RAIZ not in sys.path:
    sys.path.insert(0, _RAIZ)
from nucleo import cli  # noqa: E402
from nucleo.contrato import exito  # noqa: E402
from nucleo.registro import artefacto  # noqa: E402

BULLET_ABSTRACT_ID = 90
BULLET_NUM_ID = 90

# Fuente unica de trabajo. Las plantillas del CTEV usan Arial; el texto que
# inyectamos debe heredarla para NO mezclar tipografias en el mismo documento.
FUENTE = "Arial"
_ATRIBUTOS_FUENTE = ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia")


def _aplicar_fuente_run(run, size=None, bold=None, fuente=FUENTE):
    """Fija la fuente (en los 4 juegos de caracteres) a un run."""
    run.font.name = fuente
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for attr in _ATRIBUTOS_FUENTE:
        rf.set(qn(attr), fuente)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    return run


def _marca_fuente(parrafo, fuente=FUENTE):
    """Fija la fuente de la marca de parrafo (evita que el punto herede otra)."""
    pPr = parrafo._p.get_or_add_pPr()
    rpr = pPr.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        pPr.append(rpr)
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in _ATRIBUTOS_FUENTE:
        rf.set(qn(attr), fuente)


def forzar_fuente(document, fuente=FUENTE):
    """Fuerza una UNICA fuente en todo el documento (docDefaults + estilos clave)."""
    styles_el = document.styles.element
    dd = styles_el.find(qn("w:docDefaults"))
    if dd is not None:
        rpr_def = dd.find(qn("w:rPrDefault"))
        if rpr_def is not None:
            rpr = rpr_def.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr"); rpr_def.append(rpr)
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
            for attr in _ATRIBUTOS_FUENTE:
                rf.set(qn(attr), fuente)
    for nombre in ("Normal", "List Paragraph", "List Bullet"):
        try:
            st = document.styles[nombre]
        except KeyError:
            continue
        st.font.name = fuente
        rpr = st.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for attr in _ATRIBUTOS_FUENTE:
            rf.set(qn(attr), fuente)
    return document


def ensure_bullet_numbering(document, abstract_id=BULLET_ABSTRACT_ID, num_id=BULLET_NUM_ID):
    """Agrega una definicion de vineta (bullet) al numbering.xml si no existe."""
    numbering = document.part.numbering_part.element
    for a in numbering.findall(qn("w:abstractNum")):
        if a.get(qn("w:abstractNumId")) == str(abstract_id):
            return num_id

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    mlt = OxmlElement("w:multiLevelType"); mlt.set(qn("w:val"), "hybridMultilevel"); abstract.append(mlt)
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "bullet"); lvl.append(fmt)
    txt = OxmlElement("w:lvlText"); txt.set(qn("w:val"), "\u2022"); lvl.append(txt)
    jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); lvl.append(jc)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "720"); ind.set(qn("w:hanging"), "360"); ppr.append(ind)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), "Symbol"); rf.set(qn("w:hAnsi"), "Symbol"); rf.set(qn("w:hint"), "default")
    rpr.append(rf)
    lvl.append(rpr)
    abstract.append(lvl)
    numbering.insert(0, abstract)

    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    an = OxmlElement("w:abstractNumId"); an.set(qn("w:val"), str(abstract_id)); num.append(an)
    numbering.append(num)
    return num_id


def add_bullet(contenedor, texto, num_id=BULLET_NUM_ID, size=11, fuente=FUENTE):
    """Agrega un parrafo con vineta real."""
    p = contenedor.add_paragraph(style="List Paragraph")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id))
    numPr.append(ilvl); numPr.append(nid)
    pPr.append(numPr)
    _aplicar_fuente_run(p.add_run(texto), size=size, fuente=fuente)
    _marca_fuente(p, fuente)
    return p


def add_parrafo(contenedor, texto, size=11, fuente=FUENTE):
    p = contenedor.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _aplicar_fuente_run(p.add_run(texto), size=size, fuente=fuente)
    _marca_fuente(p, fuente)
    return p


def add_imagen(contenedor, ruta, ancho_pulgadas=6.3):
    """Inserta una imagen centrada en el contenedor."""
    p = contenedor.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(ruta, width=Inches(ancho_pulgadas))
    return p


def limpiar_celda(cell):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)


def _volcar_celda(cell, contenido, fuente=FUENTE):
    """Escribe parrafos, vinetas e imagenes dentro de una celda."""
    for txt in contenido.get("parrafos", []):
        add_parrafo(cell, txt, fuente=fuente)
    for b in contenido.get("vinetas", []):
        add_bullet(cell, b, fuente=fuente)
    for img in contenido.get("imagenes", []):
        ruta = img.get("ruta") if isinstance(img, dict) else img
        ancho = img.get("ancho", 6.3) if isinstance(img, dict) else 6.3
        add_imagen(cell, ruta, ancho)


def rellenar(plantilla, salida, encabezado, secciones, titulo="Protocolo individual",
             quitar_nota=True, fuente=FUENTE, anexo=None):
    """Rellena la tabla de una plantilla.

    encabezado: lista de lineas para la celda titulo (fila 0).
    secciones: dict {etiqueta_normalizada: {"parrafos": [...], "vinetas": [...]}}
               o lista de dicts [{"titulo","parrafos","vinetas"}] en orden.
    fuente: fuente unica a forzar en todo el documento (por defecto Arial).
    anexo: dict {"titulo": ..., "parrafos": [...], "imagenes": [...]} que se agrega
           como ULTIMA fila de la tabla (al final del trabajo).
    """
    d = docx.Document(plantilla)
    ensure_bullet_numbering(d)
    forzar_fuente(d, fuente)
    tb = d.tables[0]

    if quitar_nota:
        for par in list(d.paragraphs):
            if "Nota:" in par.text or "Instructivo" in par.text:
                par._element.getparent().remove(par._element)

    # Fila 0: titulo + encabezado
    c0 = tb.rows[0].cells[0]
    limpiar_celda(c0)
    t = c0.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _aplicar_fuente_run(t.add_run(titulo), size=16, bold=True, fuente=fuente)
    _marca_fuente(t, fuente)
    for linea in encabezado:
        p = c0.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _aplicar_fuente_run(p.add_run(linea), size=11, fuente=fuente)
        _marca_fuente(p, fuente)

    def _norm(s):
        return s.strip().rstrip(".").lower()

    if isinstance(secciones, dict):
        mapa = {_norm(k): v for k, v in secciones.items()}
    else:
        mapa = None

    for ri in range(1, len(tb.rows)):
        cell = tb.rows[ri].cells[0]
        etiqueta = cell.text.strip()
        contenido = mapa.get(_norm(etiqueta)) if mapa else (secciones[ri - 1] if ri - 1 < len(secciones) else None)
        limpiar_celda(cell)
        # conservar la etiqueta como titulo en negrita
        p0 = cell.add_paragraph(); p0.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _aplicar_fuente_run(p0.add_run(etiqueta), size=12, bold=True, fuente=fuente)
        _marca_fuente(p0, fuente)
        if not contenido:
            continue
        _volcar_celda(cell, contenido, fuente)

    if anexo:
        row = tb.add_row()
        # Evita que la fila del anexo se parta entre dos paginas
        trPr = row._tr.find(qn("w:trPr"))
        if trPr is None:
            trPr = OxmlElement("w:trPr")
            row._tr.insert(0, trPr)
        trPr.append(OxmlElement("w:cantSplit"))
        cell = row.cells[0]
        limpiar_celda(cell)
        p0 = cell.add_paragraph(); p0.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _aplicar_fuente_run(p0.add_run(anexo.get("titulo", "Anexo")), size=12, bold=True, fuente=fuente)
        _marca_fuente(p0, fuente)
        _volcar_celda(cell, anexo, fuente)

    # Quita parrafos vacios al final del cuerpo (evitan paginas en blanco)
    cuerpo = d.element.body
    for p in reversed(cuerpo.findall(qn("w:p"))):
        txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
        tiene_img = p.findall(".//" + qn("w:drawing")) or p.findall(".//" + qn("w:pict"))
        if txt.strip() or tiene_img:
            break
        cuerpo.remove(p)

    d.save(salida)
    return salida


def exportar_pdf(docx_path, pdf_path=None):
    """Exporta a PDF usando Microsoft Word (COM via PowerShell, sin dependencias)."""
    import subprocess
    pdf_path = pdf_path or os.path.splitext(docx_path)[0] + ".pdf"
    ps = (
        "$ErrorActionPreference='Stop'; "
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        "try { $d = $w.Documents.Open('%s', $false, $true); "
        "$d.SaveAs([ref]'%s', [ref]17); $d.Close($false) } finally { $w.Quit() }"
        % (docx_path, pdf_path)
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    return pdf_path


def build(plantilla, salida, encabezado, secciones, **kw):
    return rellenar(plantilla, salida, encabezado, secciones, **kw)


def _construir(ap):
    ap.add_argument("--plantilla", required=True)
    ap.add_argument("--contenido", required=True, help="JSON con {encabezado:[...], secciones:{...}}")
    ap.add_argument("--salida", required=True)
    ap.add_argument("--titulo", default="Protocolo individual")
    ap.add_argument("--pdf", action="store_true", help="Exportar tambien a PDF (Word COM)")


def _accion(ns):
    data = json.load(open(ns.contenido, encoding="utf-8"))
    rellenar(ns.plantilla, ns.salida, data.get("encabezado", []),
             data.get("secciones", {}), titulo=ns.titulo)
    arts = [artefacto("docx", ns.salida)]
    datos = {"docx": ns.salida}
    if ns.pdf:
        pdf = exportar_pdf(ns.salida)
        datos["pdf"] = pdf
        arts.append(artefacto("pdf", pdf))
    return exito(datos=datos, artefactos=arts)


if __name__ == "__main__":
    sys.exit(cli.correr("navegacion_web", _construir, _accion, sys.argv[1:],
                        prog="docx_util",
                        descripcion="Rellena una plantilla DOCX con vinetas reales."))
