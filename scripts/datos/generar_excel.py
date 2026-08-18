#!/usr/bin/env python3
"""Generador de reportes Excel profesionales a partir de JSON.

Acepta un archivo .json o JSON inline por CLI, aplica diseño institucional
(encabezados azules, bordes, ancho automático), soporta fórmulas de SUMA
o PROMEDIO por columna y formato de moneda, y emite JSON de resultado.
"""

import argparse
import importlib
import json
import os
import subprocess
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    OPENPYXL_DISPONIBLE = True
except ImportError:
    OPENPYXL_DISPONIBLE = False


def _emitir_resultado(exito, datos):
    print(json.dumps({"exito": exito, **datos}, ensure_ascii=False, indent=2))
    sys.exit(0 if exito else 1)


def _instalar_dependencia():
    """Instala openpyxl en el intérprete activo. Devuelve True si tuvo éxito."""
    try:
        resultado = subprocess.run(
            [sys.executable, "-m", "pip", "install", "openpyxl>=3.1.0"],
            capture_output=True,
            text=True,
            timeout=300,
        )
        return resultado.returncode == 0
    except (OSError, subprocess.TimeoutExpired):
        return False


def _garantizar_openpyxl(instalar_auto=False):
    if OPENPYXL_DISPONIBLE:
        return True
    if not instalar_auto:
        _emitir_resultado(
            False,
            {
                "mensaje": (
                    "Dependencia faltante: openpyxl. Instalela con "
                    "pip install openpyxl o ejecute con --instalar-deps."
                ),
                "error": "ModuleNotFoundError: openpyxl",
            },
        )
    if not _instalar_dependencia():
        _emitir_resultado(
            False,
            {
                "mensaje": "No se pudo instalar openpyxl automaticamente.",
                "error": "pip install openpyxl fallo",
            },
        )
    try:
        importlib.import_module("openpyxl")
        return True
    except ImportError:
        _emitir_resultado(
            False,
            {
                "mensaje": "openpyxl se instalo pero no es importable.",
                "error": "ImportError tras instalacion",
            },
        )
    return False


def _cargar_entrada(valor_json, archivo_json):
    """Devuelve el dict de entrada desde JSON inline o archivo."""
    origen = archivo_json or valor_json
    if origen is None:
        _emitir_resultado(
            False,
            {
                "mensaje": "Se requiere --json-data o --json-file.",
                "error": "sin entrada",
            },
        )
    if os.path.isfile(origen):
        try:
            with open(origen, encoding="utf-8") as archivo:
                return json.load(archivo)
        except (OSError, json.JSONDecodeError) as exc:
            _emitir_resultado(
                False,
                {
                    "mensaje": f"No se pudo leer el archivo de entrada: {origen}",
                    "error": str(exc),
                },
            )
    try:
        return json.loads(origen)
    except json.JSONDecodeError as exc:
        _emitir_resultado(
            False,
            {
                "mensaje": (
                    f"La entrada no es un archivo existente ni JSON valido: {origen}"
                ),
                "error": str(exc),
            },
        )
    return None


def _validar_datos(datos):
    """Valida la estructura esperada y devuelve una versión normalizada."""
    if not isinstance(datos, dict):
        raise ValueError("la entrada debe ser un objeto JSON")
    columnas = datos.get("columnas")
    filas = datos.get("filas")
    if not isinstance(columnas, list) or not columnas:
        raise ValueError("'columnas' debe ser una lista no vacia")
    if not isinstance(filas, list):
        raise ValueError("'filas' debe ser una lista de listas")
    for indice, fila in enumerate(filas):
        if not isinstance(fila, list):
            raise ValueError(f"la fila {indice + 1} no es una lista")
        if len(fila) != len(columnas):
            raise ValueError(
                f"la fila {indice + 1} tiene {len(fila)} valores y se esperaban "
                f"{len(columnas)}"
            )
    totales_raw = datos.get("totales", {})
    if isinstance(totales_raw, list):
        totales = {columna: "suma" for columna in totales_raw}
    elif isinstance(totales_raw, dict):
        totales = {k: str(v).lower() for k, v in totales_raw.items()}
    else:
        raise ValueError("'totales' debe ser un objeto o una lista de columnas")
    for columna, modo in totales.items():
        if columna not in columnas:
            raise ValueError(f"columna de totales inexistente: {columna}")
        if modo not in ("suma", "promedio"):
            raise ValueError(f"modo de total invalido para '{columna}': {modo}")
    formato = datos.get("formato_moneda", [])
    if formato is True:
        formato = [
            columna
            for indice, columna in enumerate(columnas)
            if _columna_numerica(filas, indice)
        ]
    if not isinstance(formato, list):
        raise ValueError("'formato_moneda' debe ser una lista de columnas o true")
    for columna in formato:
        if columna not in columnas:
            raise ValueError(f"columna de moneda inexistente: {columna}")
    return {
        "titulo": str(datos.get("titulo", "Reporte")),
        "columnas": [str(c) for c in columnas],
        "filas": filas,
        "totales": totales,
        "formato_moneda": list(formato),
    }


def _columna_numerica(filas, indice):
    valores = [fila[indice] for fila in filas if fila[indice] not in (None, "")]
    return bool(valores) and all(
        isinstance(v, (int, float)) and not isinstance(v, bool) for v in valores
    )


def _datos_prueba():
    """Devuelve un reporte ficticio de finanzas para --test."""
    productos = [
        ("Laptop UltraBook", "Computo", 3, 2450.00),
        ("Monitor 27 pulgadas", "Computo", 6, 780.50),
        ("Teclado mecanico", "Perifericos", 12, 145.90),
        ("Mouse inalambrico", "Perifericos", 15, 39.99),
        ("Impresora laser", "Oficina", 2, 890.00),
        ("Tablet 10 pulgadas", "Computo", 4, 560.75),
        ("Audifonos bluetooth", "Perifericos", 10, 210.25),
        ("Disco SSD 1TB", "Almacenamiento", 8, 320.00),
        ("Memoria RAM 16GB", "Almacenamiento", 9, 180.40),
        ("Webcam 4K", "Perifericos", 7, 155.30),
        ("Router Wi-Fi 6", "Redes", 5, 199.99),
        ("Cargador USB-C", "Accesorios", 20, 45.50),
        ("Funda protectora", "Accesorios", 14, 28.90),
        ("Bateria externa", "Accesorios", 11, 89.75),
        ("Cable HDMI 2.1", "Accesorios", 25, 19.99),
    ]
    columnas = ["Producto", "Categoria", "Cantidad", "Precio Unitario", "Total"]
    filas = [
        [nombre, categoria, cantidad, precio, round(cantidad * precio, 2)]
        for nombre, categoria, cantidad, precio in productos
    ]
    return {
        "titulo": "Reporte Financiero de Prueba",
        "columnas": columnas,
        "filas": filas,
        "totales": {"Cantidad": "suma", "Precio Unitario": "promedio", "Total": "suma"},
        "formato_moneda": ["Precio Unitario", "Total"],
    }


def _estilo_tabla(hoja, numero_filas, numero_columnas, fila_datos, fila_total):
    borde_delgado = Border(
        left=Side(style="thin", color="B0B0B0"),
        right=Side(style="thin", color="B0B0B0"),
        top=Side(style="thin", color="B0B0B0"),
        bottom=Side(style="thin", color="B0B0B0"),
    )
    borde_encabezado = Border(
        left=Side(style="thin", color="1F4E79"),
        right=Side(style="thin", color="1F4E79"),
        top=Side(style="thin", color="1F4E79"),
        bottom=Side(style="medium", color="1F4E79"),
    )
    relleno_encabezado = PatternFill("solid", fgColor="1F4E79")
    relleno_total = PatternFill("solid", fgColor="D9E2F3")
    fuente_encabezado = Font(bold=True, color="FFFFFF", size=11)
    fuente_total = Font(bold=True, color="1F4E79", size=11)
    for columna in range(1, numero_columnas + 1):
        celda = hoja.cell(row=fila_datos - 1, column=columna)
        celda.fill = relleno_encabezado
        celda.font = fuente_encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde_encabezado
    for fila in range(fila_datos, fila_total):
        for columna in range(1, numero_columnas + 1):
            hoja.cell(row=fila, column=columna).border = borde_delgado
    if fila_total:
        for columna in range(1, numero_columnas + 1):
            celda = hoja.cell(row=fila_total, column=columna)
            celda.fill = relleno_total
            celda.font = fuente_total
            celda.border = Border(
                left=Side(style="thin", color="1F4E79"),
                right=Side(style="thin", color="1F4E79"),
                top=Side(style="double", color="1F4E79"),
                bottom=Side(style="thin", color="1F4E79"),
            )


def _ajustar_anchos(hoja, columnas, filas, fila_datos, fila_total, numero_columnas):
    for indice, encabezado in enumerate(columnas, start=1):
        ancho_maximo = len(str(encabezado)) + 2
        for fila in filas:
            valor = fila[indice - 1]
            if valor is not None:
                ancho_maximo = max(ancho_maximo, len(str(valor)) + 2)
        if fila_total:
            celda = hoja.cell(row=fila_total, column=indice)
            ancho_maximo = max(ancho_maximo, len(str(celda.value or "")) + 2)
        hoja.column_dimensions[get_column_letter(indice)].width = min(ancho_maximo, 50)


def generar_excel(datos, ruta_salida):
    """Genera el archivo Excel y devuelve la ruta absoluta."""
    normalizado = _validar_datos(datos)
    titulo = normalizado["titulo"]
    columnas = normalizado["columnas"]
    filas = normalizado["filas"]
    totales = normalizado["totales"]
    moneda = normalizado["formato_moneda"]
    numero_columnas = len(columnas)

    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31] or "Reporte"

    fila_titulo, fila_encabezado, fila_datos = 1, 2, 3
    hoja.cell(row=fila_titulo, column=1, value=titulo)
    if numero_columnas > 1:
        hoja.merge_cells(
            start_row=fila_titulo, start_column=1, end_row=fila_titulo,
            end_column=numero_columnas,
        )
    celda_titulo = hoja.cell(row=fila_titulo, column=1)
    celda_titulo.font = Font(bold=True, size=14, color="1F4E79")
    celda_titulo.alignment = Alignment(horizontal="center", vertical="center")
    hoja.row_dimensions[fila_titulo].height = 24

    for indice, encabezado in enumerate(columnas, start=1):
        hoja.cell(row=fila_encabezado, column=indice, value=encabezado)
    for indice_fila, fila in enumerate(filas):
        for indice_columna, valor in enumerate(fila, start=1):
            hoja.cell(row=fila_datos + indice_fila, column=indice_columna, value=valor)

    fila_total = None
    if totales:
        fila_total = fila_datos + len(filas)
        hoja.cell(row=fila_total, column=1, value="TOTALES")
        ultima_datos = fila_total - 1
        for columna, modo in totales.items():
            indice = columnas.index(columna) + 1
            letra = get_column_letter(indice)
            rango = f"{letra}{fila_datos}:{letra}{ultima_datos}"
            funcion = "SUM" if modo == "suma" else "AVERAGE"
            hoja.cell(row=fila_total, column=indice, value=f"={funcion}({rango})")

    for indice, columna in enumerate(columnas):
        if columna in moneda:
            letra = get_column_letter(indice + 1)
            for fila in range(fila_datos, fila_datos + len(filas)):
                hoja.cell(row=fila, column=indice + 1).number_format = "$#,##0.00"
            if fila_total:
                hoja.cell(row=fila_total, column=indice + 1).number_format = "$#,##0.00"

    _estilo_tabla(
        hoja, len(filas), numero_columnas, fila_datos, fila_total
    )
    _ajustar_anchos(hoja, columnas, filas, fila_datos, fila_total, numero_columnas)
    hoja.freeze_panes = f"A{fila_datos}"

    ruta_absoluta = os.path.abspath(ruta_salida)
    directorio = os.path.dirname(ruta_absoluta)
    if directorio and not os.path.isdir(directorio):
        os.makedirs(directorio, exist_ok=True)
    libro.save(ruta_absoluta)
    return ruta_absoluta


def main():
    parser = argparse.ArgumentParser(
        description="Genera reportes Excel profesionales desde JSON (salida JSON)"
    )
    parser.add_argument(
        "--json-data",
        help="JSON inline o ruta a un archivo .json de entrada",
    )
    parser.add_argument(
        "--json-file",
        help="ruta explicita al archivo .json de entrada",
    )
    parser.add_argument(
        "--salida",
        default="reporte_generado.xlsx",
        help="ruta del archivo .xlsx a generar",
    )
    parser.add_argument(
        "--test",
        action="store_true",
        help="genera un reporte ficticio de finanzas de prueba",
    )
    parser.add_argument(
        "--instalar-deps",
        action="store_true",
        help="instala openpyxl automaticamente si falta",
    )
    args = parser.parse_args()

    _garantizar_openpyxl(instalar_auto=args.instalar_deps)

    try:
        if args.test:
            datos = _datos_prueba()
            if args.salida == "reporte_generado.xlsx":
                args.salida = "reporte_prueba.xlsx"
        else:
            datos = _cargar_entrada(args.json_data, args.json_file)
        ruta = generar_excel(datos, args.salida)
    except ValueError as exc:
        _emitir_resultado(
            False,
            {
                "mensaje": "No se pudo generar el archivo Excel.",
                "error": str(exc),
            },
        )
    except OSError as exc:
        _emitir_resultado(
            False,
            {
                "mensaje": "No se pudo guardar el archivo Excel.",
                "error": str(exc),
            },
        )

    filas_procesadas = len(datos["filas"])
    columnas = len(datos["columnas"])
    _emitir_resultado(
        True,
        {
            "archivo_generado": ruta,
            "filas_procesadas": filas_procesadas,
            "columnas": columnas,
            "mensaje": "Archivo Excel generado con exito.",
        },
    )


if __name__ == "__main__":
    main()
